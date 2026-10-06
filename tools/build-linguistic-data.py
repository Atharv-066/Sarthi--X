"""Build src/data/linguistic.js from the Maharashtra Linguistic Minority scheme list (.docx) + tools/linguistic-names.json.
Usage: python3 tools/build-linguistic-data.py [input.docx] [output.js]      (needs: pip install python-docx)

Layout of the Word file: one table, column 1 = language, column 2 = bullet list of scheme names. Every language gets the same
31 schemes; the five tribal languages (Bhili, Gondi, Korku, Kolami, Warli) get 16 more. The tool CHECKS the Word file against
the name table (every language, every scheme, same order) and stops with a message if anything differs, so nothing can be
dropped or reordered silently. Names are stored once (not once per language); the page builds each language from them.
"""
import json, sys, os
from docx import Document

SRC = sys.argv[1] if len(sys.argv) > 1 else 'data-source/Maharashtra_Linguistic_Minority_wise_Schemes.docx'
DST = sys.argv[2] if len(sys.argv) > 2 else 'src/data/linguistic.js'
NAMES = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'linguistic-names.json'), encoding='utf-8'))

doc = Document(SRC)
assert len(doc.tables) == 1, 'expected one table in the Word file'
rows = [(r.cells[0].text.strip(), [x.lstrip('•').strip() for x in r.cells[1].text.split('\n') if x.strip()]) for r in doc.tables[0].rows[1:]]

base, tribal, langs = NAMES['base'], NAMES['tribal'], NAMES['languages']
base_src, tribal_src = [b['src'] for b in base], [t['src'] for t in tribal]
assert [r[0] for r in rows] == [l['src'] for l in langs], 'language list in the Word file differs from linguistic-names.json'
for (name, items), lg in zip(rows, langs):
    want = base_src + (tribal_src if lg.get('tribal') else [])
    assert items == want, f'{name}: scheme list differs from the name table ({len(items)} in Word, {len(want)} expected)'
for row in base + tribal:
    assert row['sec'] in NAMES['sections'], row['src']
    assert all(row.get(k) for k in ('en', 'mr', 'hi')), 'missing name: ' + row['src']
assert len({r['en'] for r in base + tribal}) == len(base + tribal), 'two schemes share one English name'

def section(key, rows_, qual, certs):
    s = NAMES['sections'][key]
    return {'title': s['title'], 'qual': qual, 'kind': s['kind'], 'certs': certs,
            'items': [{'t': r['en']} for r in rows_ if r['sec'] == key]}

ORDER = list(NAMES['sections'])
common = [s for s in (section(k, base, '', []) for k in ORDER) if s['items']]
tribal_secs = [s for s in (section(k, tribal, 'Tribal Development Department schemes', ['st']) for k in ORDER) if s['items']]

rules = [
    'Your language only decides which schemes are shown for you. It is used for filtering and prioritising only.',
    'A scheme must still pass its own eligibility rules: religion or minority community, category, income, course, domicile, age, gender and so on.',
    'Scheme availability and eligibility can change every academic year. Confirm on the official notice before you apply.',
    'Schemes marked ST are Tribal Development Department schemes: they need the applicable Scheduled Tribe status / certificate.',
]
dictmr = {r['en']: r['mr'] for r in base + tribal if r['mr'] != r['en']}
dicthi = {r['en']: r['hi'] for r in base + tribal if r['hi'] != r['en']}
for l in langs:
    dictmr[l['name']] = l['mr']; dicthi[l['name']] = l['hi']

def dump(o): return json.dumps(o, ensure_ascii=False, separators=(',', ':'))
js = '// Linguistic-minority schemes, built from "Maharashtra Linguistic Minority wise Schemes" (.docx) + tools/linguistic-names.json by tools/build-linguistic-data.py.\n'
js += '// Scheme NAMES only (like the religious list): the source has no amounts, dates or income limits. The language only prioritises; each scheme keeps its own rules.\n'
js += '// Every language shows LING_COMMON; the five tribal languages (tribal: true) also show LING_TRIBAL. Names are stored once, not once per language.\n'
js += '// kind: scholarship | education | training | loan | development.  certs: official certificate a section needs (st).\n\n'
js += 'export const LING_COMMON = [\n' + ',\n'.join('  ' + dump(s) for s in common) + '\n];\n\n'
js += 'export const LING_TRIBAL = [\n' + ',\n'.join('  ' + dump(s) for s in tribal_secs) + '\n];\n\n'
js += 'export const LANGUAGES = [\n' + ',\n'.join('  ' + dump({k: l[k] for k in ('id', 'name') } | ({'tribal': True} if l.get('tribal') else {})) for l in langs) + '\n];\n\n'
js += 'export const LINGUISTIC = LANGUAGES.map((l) => ({ id: l.id, name: l.name, sections: l.tribal ? [...LING_COMMON, ...LING_TRIBAL] : LING_COMMON }));\n\n'
js += 'export const LING_RULES = ' + dump(rules) + ';\n\n'
js += '// English -> Marathi / Hindi names for the page translator (merged into DICT in src/translations.js).\n'
js += 'export const LING_DICT = {\n  mr: ' + dump(dictmr) + ',\n  hi: ' + dump(dicthi) + '\n};\n'
open(DST, 'w', encoding='utf-8').write(js)

print('wrote', DST, len(js) // 1024, 'KB')
print(f'{len(langs)} languages | {len(base)} common + {len(tribal)} tribal schemes | listed per language:',
      {l['name']: len(base) + (len(tribal) if l.get('tribal') else 0) for l in langs})
for s in common + tribal_secs: print('  ', s['kind'], '|', s['qual'] or '-', '|', len(s['items']))
